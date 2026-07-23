from __future__ import annotations

import csv
import io
import json
import zipfile
from datetime import UTC, datetime
from uuid import uuid4

from openpyxl import Workbook

from .db import connect
from .naming import sanitize_filename_part
from .paths import data_dir, evidence_dir, exports_dir


PLATFORM_STATUSES = {"not_assessed", "met", "partial", "not_met", "na", "escalating"}


def now_iso() -> str:
    return datetime.now(UTC).isoformat(timespec="seconds")


def new_id(prefix: str) -> str:
    return f"{prefix}_{uuid4().hex[:12]}"


def row_dict(row) -> dict:
    return dict(row) if row is not None else {}


def json_loads(value: str | None, fallback):
    if not value:
        return fallback
    try:
        return json.loads(value)
    except json.JSONDecodeError:
        return fallback


def log_event(conn, entity_type: str, entity_id: str, action: str, details: str = "") -> None:
    conn.execute(
        """
        INSERT INTO audit_events(entity_type, entity_id, action, details, created_at)
        VALUES(?, ?, ?, ?, ?)
        """,
        (entity_type, entity_id, action, details, now_iso()),
    )


def seed_default_client() -> None:
    with connect() as conn:
        existing = conn.execute("SELECT COUNT(*) FROM clients").fetchone()[0]
        if existing:
            return
        client = create_client(
            {
                "name": "Demo CMMC Client",
                "primary_contact_name": "Internal Demo",
                "primary_contact_email": "",
                "notes": "Seeded local demo client. Replace with a real customer.",
            }
        )
        update_profile(
            client["id"],
            {
                "legal_name": "Demo CMMC Client",
                "system_name": "Demo CUI Enclave",
                "environment_shape": "enclave",
                "required_cloud": "gcc",
                "current_cloud": "commercial_m365",
                "cui_type": "basic",
                "cui_flow": "CUI enters from the DoD or prime customer, is processed by authorized users, and is stored in the approved CUI workspace.",
                "cui_location": "not_started",
                "cui_users": 10,
                "timeline": "6to12",
                "internal_owner": "owner_it",
                "ongoing_support": "yes",
            },
            create_quote=True,
        )
        create_assessment(client["id"], "CMMC Level 2 Readiness")


def list_clients() -> list[dict]:
    with connect() as conn:
        rows = conn.execute(
            """
            SELECT c.*,
                   (SELECT COUNT(*) FROM assessments a WHERE a.client_id = c.id) AS assessment_count,
                   (SELECT package_name FROM quote_records q WHERE q.client_id = c.id ORDER BY q.created_at DESC LIMIT 1) AS latest_package,
                   (SELECT quote_range FROM quote_records q WHERE q.client_id = c.id ORDER BY q.created_at DESC LIMIT 1) AS latest_quote_range
            FROM clients c
            ORDER BY c.updated_at DESC, c.name
            """
        ).fetchall()
    return [row_dict(row) for row in rows]


def get_client(client_id: str) -> dict | None:
    with connect() as conn:
        client = conn.execute("SELECT * FROM clients WHERE id = ?", (client_id,)).fetchone()
        if not client:
            return None
        profile = get_profile(client_id)
        quote = latest_quote(client_id)
        assessments = conn.execute(
            "SELECT * FROM assessments WHERE client_id = ? ORDER BY created_at DESC",
            (client_id,),
        ).fetchall()
    return {
        **row_dict(client),
        "profile": profile,
        "latest_quote": quote,
        "assessments": [row_dict(row) for row in assessments],
    }


def create_client(payload: dict) -> dict:
    client_id = payload.get("id") or new_id("client")
    name = (payload.get("name") or "New Client").strip()
    ts = now_iso()
    with connect() as conn:
        conn.execute(
            """
            INSERT INTO clients(
                id, name, cage_code, uei_number, primary_contact_name,
                primary_contact_email, primary_contact_phone, notes, created_at, updated_at
            )
            VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                client_id,
                name,
                (payload.get("cage_code") or "").strip(),
                (payload.get("uei_number") or "").strip(),
                (payload.get("primary_contact_name") or "").strip(),
                (payload.get("primary_contact_email") or "").strip(),
                (payload.get("primary_contact_phone") or "").strip(),
                (payload.get("notes") or "").strip(),
                ts,
                ts,
            ),
        )
        log_event(conn, "client", client_id, "created", name)
    return get_client(client_id) or {"id": client_id, "name": name}


def update_client(client_id: str, payload: dict) -> dict | None:
    ts = now_iso()
    with connect() as conn:
        existing = conn.execute("SELECT id FROM clients WHERE id = ?", (client_id,)).fetchone()
        if not existing:
            return None
        conn.execute(
            """
            UPDATE clients
            SET name = ?, cage_code = ?, uei_number = ?, primary_contact_name = ?,
                primary_contact_email = ?, primary_contact_phone = ?, notes = ?, updated_at = ?
            WHERE id = ?
            """,
            (
                (payload.get("name") or "New Client").strip(),
                (payload.get("cage_code") or "").strip(),
                (payload.get("uei_number") or "").strip(),
                (payload.get("primary_contact_name") or "").strip(),
                (payload.get("primary_contact_email") or "").strip(),
                (payload.get("primary_contact_phone") or "").strip(),
                (payload.get("notes") or "").strip(),
                ts,
                client_id,
            ),
        )
        log_event(conn, "client", client_id, "updated")
    return get_client(client_id)


def profile_from_row(row) -> dict:
    if not row:
        return {}
    data = row_dict(row)
    data["quote_answers"] = json_loads(data.pop("quote_answers_json", "{}"), {})
    data["profile_extra"] = json_loads(data.pop("profile_json", "{}"), {})
    return data


def get_profile(client_id: str) -> dict:
    with connect() as conn:
        row = conn.execute(
            "SELECT * FROM implementation_profiles WHERE client_id = ?",
            (client_id,),
        ).fetchone()
    return profile_from_row(row)


def update_profile(client_id: str, payload: dict, create_quote: bool = False) -> dict:
    ts = now_iso()
    quote_answers = payload.get("quote_answers") or payload.get("quote_answers_json") or {}
    if isinstance(quote_answers, str):
        quote_answers = json_loads(quote_answers, {})
    profile_extra = payload.get("profile_extra") or payload.get("profile_json") or {}
    if isinstance(profile_extra, str):
        profile_extra = json_loads(profile_extra, {})

    fields = {
        "legal_name": payload.get("legal_name", ""),
        "system_name": payload.get("system_name", ""),
        "environment_shape": payload.get("environment_shape", ""),
        "required_cloud": payload.get("required_cloud", ""),
        "current_cloud": payload.get("current_cloud", ""),
        "cui_type": payload.get("cui_type", ""),
        "cui_flow": payload.get("cui_flow", ""),
        "cui_location": payload.get("cui_location", ""),
        "cui_users": normalize_user_count(payload.get("cui_users")) or None,
        "total_employees": payload.get("total_employees"),
        "total_endpoints": payload.get("total_endpoints"),
        "endpoint_management": payload.get("endpoint_management", ""),
        "mfa_status": payload.get("mfa_status", ""),
        "logging_status": payload.get("logging_status", ""),
        "encryption_status": payload.get("encryption_status", ""),
        "external_access": payload.get("external_access", ""),
        "external_service_providers": payload.get("external_service_providers", ""),
        "timeline": payload.get("timeline", ""),
        "internal_owner": payload.get("internal_owner", ""),
        "ongoing_support": payload.get("ongoing_support", ""),
        "quote_answers_json": json.dumps(quote_answers),
        "profile_json": json.dumps(profile_extra),
        "questionnaire_complete": 1 if payload.get("questionnaire_complete") else 0,
        "updated_at": ts,
    }
    with connect() as conn:
        exists = conn.execute(
            "SELECT client_id FROM implementation_profiles WHERE client_id = ?",
            (client_id,),
        ).fetchone()
        if exists:
            assignments = ", ".join(f"{key} = :{key}" for key in fields)
            conn.execute(
                f"UPDATE implementation_profiles SET {assignments} WHERE client_id = :client_id",
                {**fields, "client_id": client_id},
            )
        else:
            columns = ["client_id", *fields.keys()]
            placeholders = ", ".join(f":{column}" for column in columns)
            conn.execute(
                f"INSERT INTO implementation_profiles({', '.join(columns)}) VALUES({placeholders})",
                {**fields, "client_id": client_id},
            )
        conn.execute("UPDATE clients SET updated_at = ? WHERE id = ?", (ts, client_id))
        log_event(conn, "client", client_id, "profile_updated")
    profile = get_profile(client_id)
    if create_quote:
        create_quote_record(client_id, profile)
    return profile


def score_profile(profile: dict) -> int:
    checks = [
        bool(profile.get("cui_type") and profile.get("cui_type") != "unknown"),
        bool(profile.get("environment_shape") and profile.get("environment_shape") != "unknown"),
        bool(profile.get("required_cloud") and profile.get("required_cloud") != "unknown"),
        bool(profile.get("cui_flow")),
        bool(normalize_user_count(profile.get("cui_users"))),
        profile.get("mfa_status") in {"yes", "implemented", "enforced"},
        profile.get("endpoint_management") in {"yes", "managed"},
        profile.get("logging_status") in {"yes", "managed"},
        profile.get("encryption_status") in {"yes", "implemented"},
        bool(profile.get("internal_owner") and profile.get("internal_owner") != "none"),
    ]
    return round(sum(1 for check in checks if check) / len(checks) * 100)


def determine_quote(profile: dict) -> dict:
    cui_type = profile.get("cui_type") or ""
    shape = profile.get("environment_shape") or ""
    required_cloud = profile.get("required_cloud") or ""
    current_cloud = profile.get("current_cloud") or ""
    external_access = profile.get("external_access") or ""
    cui_users = normalize_user_count(profile.get("cui_users"))
    timeline = profile.get("timeline") or ""

    high_required = required_cloud == "gcc_high" or cui_type == "specified" or current_cloud == "gcc_high"
    controlled = required_cloud == "controlled_access" or external_access == "subs_byod"
    full_tenant = shape == "full_tenant" or (shape != "enclave" and cui_users >= 51)

    if cui_type == "fci":
        package, quote_range = "Level 1 / FCI Starter", "$0-$15k"
    elif controlled:
        package, quote_range = "Controlled Access / VDI CUI Environment + CMMC L2 Prep", "$95k-$250k+"
    elif high_required and full_tenant:
        package, quote_range = "Full GCC High Migration + CMMC L2 Prep", "$120k-$250k+"
    elif high_required:
        package, quote_range = "GCC High Enclave + CMMC L2 Prep", "$70k-$150k"
    elif full_tenant:
        package, quote_range = "Full GCC Tenant Migration + CMMC L2 Prep", "$60k-$120k"
    elif current_cloud == "gcc":
        package, quote_range = "Existing GCC Hardening + L2 Prep", "$30k-$65k"
    else:
        package, quote_range = "GCC Enclave + CMMC L2 Prep", "$40k-$75k"

    assumptions = []
    unknowns = [
        label
        for label, value in [
            ("CUI type", cui_type),
            ("environment shape", shape),
            ("required cloud", required_cloud),
            ("current cloud", current_cloud),
            ("CUI users", cui_users),
        ]
        if value in ("", None, "unknown", 0)
    ]
    if unknowns:
        assumptions.append("Confirm " + ", ".join(unknowns) + " before final SOW.")
    if timeline == "lt6":
        assumptions.append("Accelerated timeline may move the estimate toward the upper end.")
    if profile.get("ongoing_support") in {"yes", "maybe"}:
        assumptions.append("Recommend vCISO + Ongoing Compliance at $6,000/month.")
    confidence = "Low" if len(unknowns) >= 3 else "Medium" if unknowns else "High"
    return {
        "readiness_score": score_profile(profile),
        "package_name": package,
        "quote_range": quote_range,
        "confidence": confidence,
        "assumptions": assumptions,
    }


def normalize_user_count(value) -> int:
    if value in ("", None, "unknown"):
        return 0
    if isinstance(value, int):
        return value
    text = str(value).strip().lower()
    if text in {"1_10", "1-10"}:
        return 10
    if text in {"11_50", "11-50"}:
        return 50
    if text in {"51_150", "51-150"}:
        return 150
    if text in {"151_plus", "151+", "151 plus"}:
        return 151
    match = re_digits(text)
    return int(match) if match else 0


def re_digits(text: str) -> str:
    return "".join(ch for ch in text if ch.isdigit())


def create_quote_record(client_id: str, profile: dict | None = None) -> dict:
    profile = profile or get_profile(client_id)
    quote = determine_quote(profile)
    quote_id = new_id("quote")
    ts = now_iso()
    with connect() as conn:
        conn.execute(
            """
            INSERT INTO quote_records(
                id, client_id, readiness_score, package_name, quote_range,
                confidence, assumptions, answers_json, created_at
            )
            VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                quote_id,
                client_id,
                quote["readiness_score"],
                quote["package_name"],
                quote["quote_range"],
                quote["confidence"],
                "\n".join(quote["assumptions"]),
                json.dumps(profile),
                ts,
            ),
        )
        log_event(conn, "quote", quote_id, "created", quote["package_name"])
    return latest_quote(client_id) or {"id": quote_id, **quote}


def latest_quote(client_id: str) -> dict | None:
    with connect() as conn:
        row = conn.execute(
            "SELECT * FROM quote_records WHERE client_id = ? ORDER BY created_at DESC LIMIT 1",
            (client_id,),
        ).fetchone()
    return row_dict(row) if row else None


def create_assessment(client_id: str, name: str, description: str = "") -> dict:
    assessment_id = new_id("assessment")
    ts = now_iso()
    with connect() as conn:
        conn.execute(
            """
            INSERT INTO assessments(id, client_id, name, description, created_at, updated_at)
            VALUES(?, ?, ?, ?, ?, ?)
            """,
            (assessment_id, client_id, name.strip() or "CMMC Level 2 Readiness", description.strip(), ts, ts),
        )
        ensure_assessment_results(conn, assessment_id)
        log_event(conn, "assessment", assessment_id, "created", client_id)
    return get_assessment(assessment_id) or {"id": assessment_id, "client_id": client_id, "name": name}


def ensure_assessment_results(conn, assessment_id: str) -> None:
    ts = now_iso()
    objective_ids = [row["id"] for row in conn.execute("SELECT id FROM objectives ORDER BY id").fetchall()]
    for objective_id in objective_ids:
        result_id = f"{assessment_id}_{objective_id}"
        conn.execute(
            """
            INSERT OR IGNORE INTO assessment_results(id, assessment_id, objective_id, updated_at)
            VALUES(?, ?, ?, ?)
            """,
            (result_id, assessment_id, objective_id, ts),
        )


def list_assessments(client_id: str | None = None) -> list[dict]:
    sql = """
        SELECT a.*, c.name AS client_name
        FROM assessments a
        JOIN clients c ON c.id = a.client_id
    """
    params: list[str] = []
    if client_id:
        sql += " WHERE a.client_id = ?"
        params.append(client_id)
    sql += " ORDER BY a.created_at DESC"
    with connect() as conn:
        rows = conn.execute(sql, params).fetchall()
    return [row_dict(row) for row in rows]


def get_assessment(assessment_id: str) -> dict | None:
    with connect() as conn:
        assessment = conn.execute(
            """
            SELECT a.*, c.name AS client_name
            FROM assessments a
            JOIN clients c ON c.id = a.client_id
            WHERE a.id = ?
            """,
            (assessment_id,),
        ).fetchone()
        if not assessment:
            return None
        score = assessment_score(assessment_id)
    return {**row_dict(assessment), "score": score}


def assessment_score(assessment_id: str) -> dict:
    with connect() as conn:
        rows = conn.execute(
            """
            SELECT status, COUNT(*) AS count
            FROM assessment_results
            WHERE assessment_id = ?
            GROUP BY status
            """,
            (assessment_id,),
        ).fetchall()
        evidence_count = conn.execute(
            """
            SELECT COUNT(DISTINCT ae.evidence_id)
            FROM assessment_evidence ae
            JOIN assessment_results ar ON ar.id = ae.result_id
            WHERE ar.assessment_id = ?
            """,
            (assessment_id,),
        ).fetchone()[0]
    counts = {row["status"]: row["count"] for row in rows}
    total = sum(counts.values())
    met = counts.get("met", 0)
    partial = counts.get("partial", 0)
    na = counts.get("na", 0)
    denominator = max(total - na, 1)
    score = round(((met + partial * 0.5) / denominator) * 100, 1)
    return {"total": total, "score": score, "counts": counts, "evidence_count": evidence_count}


def list_controls(assessment_id: str, family: str = "", status: str = "", missing: bool = False, q: str = "") -> list[dict]:
    filters = ["ar.assessment_id = ?"]
    params: list = [assessment_id]
    if family:
        filters.append("d.code = ?")
        params.append(family)
    if status:
        filters.append("ar.status = ?")
        params.append(status)
    if q:
        filters.append("(o.id LIKE ? OR r.id LIKE ? OR r.name LIKE ? OR o.text LIKE ?)")
        term = f"%{q}%"
        params.extend([term, term, term, term])
    where = " AND ".join(filters)
    having = "HAVING evidence_count = 0" if missing else ""
    with connect() as conn:
        rows = conn.execute(
            f"""
            SELECT
                ar.id AS result_id,
                ar.status,
                ar.notes,
                ar.owner,
                ar.due_date,
                d.code AS family,
                d.name AS family_name,
                d.sort_order AS family_order,
                r.id AS requirement_id,
                r.name AS requirement_name,
                r.text AS requirement_text,
                r.potential_methods,
                r.discussion,
                r.further_discussion,
                o.id AS objective_id,
                o.text AS objective_text,
                o.evidence_examples,
                COUNT(ae.evidence_id) AS evidence_count
            FROM assessment_results ar
            JOIN objectives o ON o.id = ar.objective_id
            JOIN requirements r ON r.id = o.requirement_id
            JOIN domains d ON d.code = r.domain_code
            LEFT JOIN assessment_evidence ae ON ae.result_id = ar.id
            WHERE {where}
            GROUP BY ar.id
            {having}
            ORDER BY d.sort_order, r.id, o.letter
            """,
            params,
        ).fetchall()
    return [objective_payload(row) for row in rows]


def objective_payload(row) -> dict:
    data = row_dict(row)
    data["evidence_examples"] = json_loads(data.get("evidence_examples"), [])
    return data


def get_result(result_id: str) -> dict | None:
    with connect() as conn:
        row = conn.execute(
            """
            SELECT
                ar.id AS result_id,
                ar.assessment_id,
                ar.status,
                ar.notes,
                ar.owner,
                ar.due_date,
                d.code AS family,
                d.name AS family_name,
                r.id AS requirement_id,
                r.name AS requirement_name,
                r.text AS requirement_text,
                r.potential_methods,
                r.discussion,
                r.further_discussion,
                o.id AS objective_id,
                o.text AS objective_text,
                o.evidence_examples
            FROM assessment_results ar
            JOIN objectives o ON o.id = ar.objective_id
            JOIN requirements r ON r.id = o.requirement_id
            JOIN domains d ON d.code = r.domain_code
            WHERE ar.id = ?
            """,
            (result_id,),
        ).fetchone()
        if not row:
            return None
        evidence = conn.execute(
            """
            SELECT e.*, ae.linked_at
            FROM assessment_evidence ae
            JOIN evidence e ON e.id = ae.evidence_id
            WHERE ae.result_id = ?
            ORDER BY e.title
            """,
            (result_id,),
        ).fetchall()
    payload = objective_payload(row)
    payload["evidence"] = [row_dict(item) for item in evidence]
    return payload


def update_result(result_id: str, payload: dict) -> dict | None:
    status = payload.get("status") or "not_assessed"
    if status not in PLATFORM_STATUSES:
        status = "not_assessed"
    with connect() as conn:
        existing = conn.execute("SELECT assessment_id FROM assessment_results WHERE id = ?", (result_id,)).fetchone()
        if not existing:
            return None
        conn.execute(
            """
            UPDATE assessment_results
            SET status = ?, notes = ?, owner = ?, due_date = ?, updated_at = ?
            WHERE id = ?
            """,
            (
                status,
                (payload.get("notes") or "").strip(),
                (payload.get("owner") or "").strip(),
                (payload.get("due_date") or "").strip(),
                now_iso(),
                result_id,
            ),
        )
        sync_poam_for_result(conn, result_id, status, payload)
        log_event(conn, "result", result_id, "updated", status)
    return get_result(result_id)


def sync_poam_for_result(conn, result_id: str, status: str, payload: dict) -> None:
    row = conn.execute(
        "SELECT assessment_id, objective_id FROM assessment_results WHERE id = ?",
        (result_id,),
    ).fetchone()
    if not row:
        return
    if status not in {"not_met", "partial", "escalating"}:
        return
    existing = conn.execute(
        "SELECT id FROM poam_items WHERE assessment_id = ? AND objective_id = ? AND status <> 'closed'",
        (row["assessment_id"], row["objective_id"]),
    ).fetchone()
    if existing:
        return
    poam_id = new_id("poam")
    ts = now_iso()
    conn.execute(
        """
        INSERT INTO poam_items(
            id, assessment_id, objective_id, title, gap, remediation, owner,
            due_date, status, priority, created_at, updated_at
        )
        VALUES(?, ?, ?, ?, ?, ?, ?, ?, 'open', 'medium', ?, ?)
        """,
        (
            poam_id,
            row["assessment_id"],
            row["objective_id"],
            f"Remediate {row['objective_id']}",
            (payload.get("notes") or "").strip(),
            "Capture evidence, implement missing control activity, and update SSP/POA&M.",
            (payload.get("owner") or "").strip(),
            (payload.get("due_date") or "").strip(),
            ts,
            ts,
        ),
    )


def list_evidence_library(assessment_id: str | None = None) -> list[dict]:
    with connect() as conn:
        if assessment_id:
            rows = conn.execute(
                """
                SELECT e.*, COUNT(ar.id) AS mapped_count
                FROM evidence e
                LEFT JOIN assessment_evidence ae ON ae.evidence_id = e.id
                LEFT JOIN assessment_results ar ON ar.id = ae.result_id AND ar.assessment_id = ?
                GROUP BY e.id
                ORDER BY e.uploaded_at DESC
                """,
                (assessment_id,),
            ).fetchall()
        else:
            rows = conn.execute(
                """
                SELECT e.*, COUNT(ae.result_id) AS mapped_count
                FROM evidence e
                LEFT JOIN assessment_evidence ae ON ae.evidence_id = e.id
                GROUP BY e.id
                ORDER BY e.uploaded_at DESC
                """
            ).fetchall()
    return [row_dict(row) for row in rows]


def link_evidence(result_id: str, evidence_id: int) -> None:
    with connect() as conn:
        conn.execute(
            "INSERT OR IGNORE INTO assessment_evidence(result_id, evidence_id, linked_at) VALUES(?, ?, ?)",
            (result_id, evidence_id, now_iso()),
        )
        conn.execute(
            "UPDATE assessment_results SET status = CASE WHEN status = 'not_assessed' THEN 'partial' ELSE status END, updated_at = ? WHERE id = ?",
            (now_iso(), result_id),
        )
        log_event(conn, "result", result_id, "evidence_linked", str(evidence_id))


def unlink_evidence(result_id: str, evidence_id: int) -> None:
    with connect() as conn:
        conn.execute(
            "DELETE FROM assessment_evidence WHERE result_id = ? AND evidence_id = ?",
            (result_id, evidence_id),
        )
        log_event(conn, "result", result_id, "evidence_unlinked", str(evidence_id))


def tailored_evidence_guidance(client_id: str, objective_id: str) -> dict:
    profile = get_profile(client_id)
    with connect() as conn:
        row = conn.execute(
            """
            SELECT o.id AS objective_id, o.text AS objective_text, o.evidence_examples,
                   r.id AS requirement_id, r.name AS requirement_name,
                   r.potential_methods, r.discussion, d.name AS family_name
            FROM objectives o
            JOIN requirements r ON r.id = o.requirement_id
            JOIN domains d ON d.code = r.domain_code
            WHERE o.id = ?
            """,
            (objective_id,),
        ).fetchone()
    if not row:
        return {}
    examples = json_loads(row["evidence_examples"], [])
    tools = []
    if profile.get("required_cloud") in {"gcc", "gcc_high"} or profile.get("current_cloud") in {"gcc", "gcc_high", "commercial_m365"}:
        tools.extend(["Microsoft Entra ID", "Microsoft Intune", "Microsoft Purview", "Microsoft Defender"])
    if profile.get("endpoint_management") in {"yes", "managed"}:
        tools.append("EDR / endpoint management console")
    if profile.get("logging_status") in {"yes", "managed"}:
        tools.append("SIEM / audit log platform")
    if not tools:
        tools.append("system owner interview and configuration screenshots")
    guidance = [
        f"Confirm how {row['requirement_id']} applies inside {profile.get('system_name') or 'the CUI environment'}.",
        "Capture screenshots or exports that show the configured control, scope, date, and tenant/system name.",
        "Prefer authoritative admin portals, approved policies/procedures, tickets, reports, and logs over informal notes.",
    ]
    if profile.get("external_access"):
        guidance.append("Validate whether subcontractor/MSP access is inside the approved CUI boundary.")
    return {
        "objective_id": row["objective_id"],
        "requirement_id": row["requirement_id"],
        "family_name": row["family_name"],
        "evidence_examples": examples,
        "recommended_sources": tools,
        "guidance": guidance,
        "official_discussion": row["discussion"],
        "potential_methods": row["potential_methods"],
    }


def render_template_body(body: str, client: dict, profile: dict, assessment: dict | None = None) -> str:
    quote = latest_quote(client["id"]) or {}
    replacements = {
        "client_name": client.get("name") or "",
        "system_name": profile.get("system_name") or "CUI Environment",
        "required_cloud": fmt_choice(profile.get("required_cloud")),
        "environment_shape": fmt_choice(profile.get("environment_shape")),
        "external_access": fmt_choice(profile.get("external_access")),
        "cui_summary": profile.get("cui_flow") or "CUI flow has not been fully documented yet.",
        "environment_summary": environment_summary(profile, quote),
        "control_summary": control_summary_text(assessment["id"]) if assessment else "No assessment selected.",
        "evidence_summary": evidence_summary_text(assessment["id"]) if assessment else "No assessment selected.",
    }
    rendered = body
    for key, value in replacements.items():
        rendered = rendered.replace("{{" + key + "}}", str(value))
    return rendered


def fmt_choice(value: str | None) -> str:
    return (value or "Not specified").replace("_", " ").title()


def environment_summary(profile: dict, quote: dict) -> str:
    return (
        f"Target package: {quote.get('package_name', 'Not quoted yet')}\n"
        f"Quote range: {quote.get('quote_range', 'Not quoted yet')}\n"
        f"CUI users: {profile.get('cui_users') or 'Not specified'}\n"
        f"Current cloud: {fmt_choice(profile.get('current_cloud'))}\n"
        f"Endpoint management: {fmt_choice(profile.get('endpoint_management'))}\n"
        f"External access: {fmt_choice(profile.get('external_access'))}"
    )


def control_summary_text(assessment_id: str) -> str:
    score = assessment_score(assessment_id)
    counts = score["counts"]
    return (
        f"Total objectives: {score['total']}\n"
        f"Score: {score['score']}%\n"
        f"Met: {counts.get('met', 0)}\n"
        f"Partial: {counts.get('partial', 0)}\n"
        f"Not met: {counts.get('not_met', 0)}\n"
        f"Escalating: {counts.get('escalating', 0)}"
    )


def evidence_summary_text(assessment_id: str) -> str:
    score = assessment_score(assessment_id)
    return f"Distinct evidence files mapped: {score['evidence_count']}"


def generate_documents(client_id: str, assessment_id: str | None = None) -> list[dict]:
    client = get_client(client_id)
    if not client:
        return []
    profile = client.get("profile") or {}
    assessment = get_assessment(assessment_id) if assessment_id else None
    ts = now_iso()
    generated: list[dict] = []
    with connect() as conn:
        templates = conn.execute("SELECT * FROM document_templates ORDER BY doc_type, title").fetchall()
        for template in templates:
            document_id = new_id("doc")
            body = render_template_body(template["body"], client, profile, assessment)
            title = f"{client['name']} - {template['title']}"
            filename = sanitize_filename_part(title) + ".md"
            conn.execute(
                """
                INSERT INTO generated_documents(
                    id, client_id, assessment_id, doc_type, title, body,
                    filename, created_at, updated_at
                )
                VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    document_id,
                    client_id,
                    assessment_id,
                    template["doc_type"],
                    title,
                    body,
                    filename,
                    ts,
                    ts,
                ),
            )
            generated.append(
                {
                    "id": document_id,
                    "client_id": client_id,
                    "assessment_id": assessment_id,
                    "doc_type": template["doc_type"],
                    "title": title,
                    "filename": filename,
                }
            )
        log_event(conn, "client", client_id, "documents_generated", str(len(generated)))
    return generated


def list_documents(client_id: str, assessment_id: str | None = None) -> list[dict]:
    sql = "SELECT * FROM generated_documents WHERE client_id = ?"
    params: list = [client_id]
    if assessment_id:
        sql += " AND assessment_id = ?"
        params.append(assessment_id)
    sql += " ORDER BY created_at DESC"
    with connect() as conn:
        rows = conn.execute(sql, params).fetchall()
    return [row_dict(row) for row in rows]


def save_document(document_id: str, payload: dict) -> dict | None:
    with connect() as conn:
        existing = conn.execute("SELECT * FROM generated_documents WHERE id = ?", (document_id,)).fetchone()
        if not existing:
            return None
        conn.execute(
            """
            UPDATE generated_documents
            SET title = ?, body = ?, updated_at = ?
            WHERE id = ?
            """,
            (
                (payload.get("title") or existing["title"]).strip(),
                payload.get("body") or "",
                now_iso(),
                document_id,
            ),
        )
    with connect() as conn:
        row = conn.execute("SELECT * FROM generated_documents WHERE id = ?", (document_id,)).fetchone()
    return row_dict(row)


def poam_rows(assessment_id: str) -> list[dict]:
    with connect() as conn:
        rows = conn.execute(
            """
            SELECT p.*, r.name AS requirement_name, d.name AS family_name
            FROM poam_items p
            JOIN objectives o ON o.id = p.objective_id
            JOIN requirements r ON r.id = o.requirement_id
            JOIN domains d ON d.code = r.domain_code
            WHERE p.assessment_id = ?
            ORDER BY p.status, p.priority, p.due_date, p.objective_id
            """,
            (assessment_id,),
        ).fetchall()
    return [row_dict(row) for row in rows]


def evidence_capture_rows(assessment_id: str) -> list[dict]:
    with connect() as conn:
        rows = conn.execute(
            """
            SELECT
                d.name AS domain,
                r.id AS requirement_id,
                r.name AS requirement_name,
                o.id AS objective_id,
                o.text AS objective_text,
                ar.status,
                ar.notes,
                COUNT(ae.evidence_id) AS evidence_count
            FROM assessment_results ar
            JOIN objectives o ON o.id = ar.objective_id
            JOIN requirements r ON r.id = o.requirement_id
            JOIN domains d ON d.code = r.domain_code
            LEFT JOIN assessment_evidence ae ON ae.result_id = ar.id
            WHERE ar.assessment_id = ?
            GROUP BY ar.id
            ORDER BY d.sort_order, r.id, o.letter
            """,
            (assessment_id,),
        ).fetchall()
    return [row_dict(row) for row in rows]


def rows_to_xlsx(sheet_name: str, headers: list[str], rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]
    ws.append(headers)
    for row in rows:
        ws.append([row.get(header, "") for header in headers])
    for column_cells in ws.columns:
        width = min(max(len(str(cell.value or "")) for cell in column_cells) + 2, 70)
        ws.column_dimensions[column_cells[0].column_letter].width = width
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def evidence_report_xlsx(assessment_id: str) -> bytes:
    headers = ["domain", "requirement_id", "requirement_name", "objective_id", "objective_text", "status", "evidence_count", "notes"]
    return rows_to_xlsx("Evidence Capture", headers, evidence_capture_rows(assessment_id))


def poam_report_xlsx(assessment_id: str) -> bytes:
    headers = ["family_name", "objective_id", "title", "gap", "remediation", "owner", "due_date", "status", "priority"]
    return rows_to_xlsx("POAM", headers, poam_rows(assessment_id))


def export_platform_package(client_id: str, assessment_id: str) -> Path:
    client = get_client(client_id)
    assessment = get_assessment(assessment_id)
    if not client or not assessment:
        raise ValueError("Client or assessment not found")
    exports_dir().mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    base = sanitize_filename_part(client["name"])
    zip_path = exports_dir() / f"{base}_cmmc_package_{stamp}.zip"
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("reports/evidence_capture.xlsx", evidence_report_xlsx(assessment_id))
        archive.writestr("reports/poam.xlsx", poam_report_xlsx(assessment_id))
        archive.writestr("profile/client_profile.json", json.dumps(client, indent=2, default=str))
        for document in list_documents(client_id, assessment_id):
            folder = document["doc_type"] + "s"
            archive.writestr(f"documents/{folder}/{document['filename']}", document["body"])
        write_evidence_files(archive, assessment_id)
    return zip_path


def write_evidence_files(archive: zipfile.ZipFile, assessment_id: str) -> None:
    manifest = io.StringIO(newline="")
    writer = csv.writer(manifest)
    writer.writerow(["Objective ID", "Evidence ID", "Evidence Title", "Original Filename", "Export Path", "SHA256", "Capture Date", "Notes"])
    used_names: set[str] = set()
    with connect() as conn:
        rows = conn.execute(
            """
            SELECT o.id AS objective_id, e.*
            FROM assessment_evidence ae
            JOIN assessment_results ar ON ar.id = ae.result_id
            JOIN objectives o ON o.id = ar.objective_id
            JOIN evidence e ON e.id = ae.evidence_id
            WHERE ar.assessment_id = ?
            ORDER BY o.id, e.title
            """,
            (assessment_id,),
        ).fetchall()
    for row in rows:
        safe_title = sanitize_filename_part(row["title"])
        filename = f"{row['objective_id']}-{safe_title}{row['extension']}"
        counter = 2
        while filename in used_names:
            filename = f"{row['objective_id']}-{safe_title}-{counter}{row['extension']}"
            counter += 1
        used_names.add(filename)
        export_path = f"evidence/{filename}"
        source = evidence_dir().parent / row["stored_path"]
        if source.exists():
            archive.write(source, export_path)
        writer.writerow(
            [
                row["objective_id"],
                row["id"],
                row["title"],
                row["original_filename"],
                export_path,
                row["sha256"],
                row["capture_date"],
                row["notes"],
            ]
        )
    archive.writestr("evidence/manifest.csv", manifest.getvalue())
